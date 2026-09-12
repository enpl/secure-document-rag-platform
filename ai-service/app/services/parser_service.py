"""F-AI-003. PDF/DOCX/XLSX/일반 텍스트 계열 텍스트 추출(M06 Content Processing
Core, 승인된 초기 지원 集合 - CLAUDE.md Multi-format Source Model의 최종
목표가 아니라, 이 작업이 실제로 구현/검증하는 좁은 초기 集合이다).

이 파일이 유일한 Parser 구현이다(Java 쪽 병렬 구현 없음). Backend는 이
Service를 HTTP로만 호출한다({@code DocumentParsingClient}).

핵심 보안 경계:
  - 형식은 확장자/선언된 MIME/실제 Byte Signature/Container 구조를 함께
    본다 - 하나만으로 판단하지 않는다(Spoofing 방어).
  - 분류(Classification)와 Container 검사 자체도 신뢰할 수 없는 입력에 대한
    처리이므로, 실제 Parsing과 동일하게 Timeout/동시성 상한이 적용되는
    자식 Process 안에서 수행한다(M06 후속 교정 - 이전에는 부모 Process에서
    수행되어 이 경계 밖에 있었다).
  - Zip 기반 Container(DOCX/XLSX)는 Central Directory의 선언된 크기뿐 아니라
    실제로 읽은 압축 해제 Byte 수도 함께 검증한다(선언된 크기를 위조한
    Archive 방어) - 개별 Entry 비율과 전체 집계 비율을 모두 본다.
  - Path Traversal Entry 이름은 거부한다.
  - Macro/Active Content(vbaProject.bin)가 포함된 Container는 "우리가
    실행하지 않는다"로 끝내지 않고 적극적으로 거부한다.
  - 외부 관계/네트워크 자원 로딩은 시도하지 않는다(모든 Parsing은 완전히
    오프라인).
  - 중첩 Archive는 재귀적으로 풀지 않는다(Core MVP 제외 사항).
  - 실제 실행 시간이 설정된 Timeout을 넘기면 그 작업을 수행 중인 별도
    Process를 강제 종료해 자원을 회수한다 - Queue로 결과를 받을 때
    `process.join(timeout)`을 먼저 호출하지 않는다(아래 `_run_with_timeout`
    Docstring의 Deadlock 설명 참고).
  - 동시 활성 Parser 작업 수는 Non-Blocking으로 제한한다(Bounded Admission) -
    한도를 넘으면 대기열에 쌓지 않고 즉시 실패로 응답한다.
  - 출력 길이 상한은 추출 도중(전체를 다 모으기 전에) 확인해 조기 종료한다.
  - 실패하는 입력의 원문 내용은 로그에 남기지 않는다.
"""

from __future__ import annotations

import io
import multiprocessing as mp
import queue as queue_module
import threading
import time
import zipfile
from typing import List, Optional, Tuple

from app.core.config import Settings
from app.models.schemas import (
    OUTCOME_FAILED,
    OUTCOME_NO_TEXT,
    OUTCOME_SUCCESS,
    OUTCOME_UNSUPPORTED_FORMAT,
    LocationDto,
    ParseResponseDto,
)

NORMALIZATION_VERSION = "1"

# ---------------------------------------------------------------------------
# 승인된 초기 지원 集合 (M06 구현 계약)
# ---------------------------------------------------------------------------
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".java", ".py", ".js", ".ts", ".sql",
}
_PDF_EXTENSIONS = {".pdf"}
_DOCX_EXTENSIONS = {".docx"}
_XLSX_EXTENSIONS = {".xlsx"}

_FAMILY_TEXT = "TEXT"
_FAMILY_PDF = "PDF"
_FAMILY_DOCX = "DOCX"
_FAMILY_XLSX = "XLSX"

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"


def _extension_of(file_name: str) -> str:
    dot = file_name.rfind(".")
    return file_name[dot:].lower() if dot >= 0 else ""


def _extension_family(file_name: str) -> Optional[str]:
    ext = _extension_of(file_name)
    if ext in _TEXT_EXTENSIONS:
        return _FAMILY_TEXT
    if ext in _PDF_EXTENSIONS:
        return _FAMILY_PDF
    if ext in _DOCX_EXTENSIONS:
        return _FAMILY_DOCX
    if ext in _XLSX_EXTENSIONS:
        return _FAMILY_XLSX
    return None


def _declared_mime_family(declared_mime_type: Optional[str]) -> Optional[str]:
    """선언된 MIME이 명확히 어떤 family를 주장하면 그 family를, 모호하거나
    없으면 None(검사에서 제외 - 증거로 취급하지 않는다)을 반환한다."""
    if not declared_mime_type:
        return None
    value = declared_mime_type.split(";")[0].strip().lower()
    if value == "application/pdf":
        return _FAMILY_PDF
    if value == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _FAMILY_DOCX
    if value == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _FAMILY_XLSX
    if value.startswith("text/") or value in ("application/json", "application/xml", "application/yaml"):
        return _FAMILY_TEXT
    return None


def _detected_family(content: bytes) -> Optional[str]:
    """실제 Byte로만 판단한다(확장자/선언된 MIME과 무관)."""
    if content.startswith(_PDF_MAGIC):
        return _FAMILY_PDF
    if content.startswith(_ZIP_MAGIC):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile:
            return None
        if "word/document.xml" in names:
            return _FAMILY_DOCX
        if "xl/workbook.xml" in names:
            return _FAMILY_XLSX
        return None
    # 일반 텍스트 계열은 고유 Magic Byte가 없다 - 디코딩 가능 여부로 판단한다.
    if _looks_like_text(content):
        return _FAMILY_TEXT
    return None


def _looks_like_text(content: bytes) -> bool:
    if b"\x00" in content:
        return False
    try:
        from charset_normalizer import from_bytes

        result = from_bytes(content).best()
        return result is not None
    except Exception:
        return False


def classify_format(content: bytes, file_name: str, declared_mime_type: Optional[str]) -> Optional[str]:
    """확장자/선언된 MIME/실제 Byte Signature가 모두 같은 family를 가리킬
    때만 그 family를 반환한다 - 하나라도 다른 family를 명확히 주장하면
    Spoofing 의심으로 None(미지원/거부)을 반환한다. 확장자가 애초에 알려진
    지원 集合 밖이면(예: .pptx) 그 자체로 None이다."""
    extension_family = _extension_family(file_name)
    if extension_family is None:
        return None
    detected = _detected_family(content)
    if detected is None or detected != extension_family:
        return None
    declared = _declared_mime_family(declared_mime_type)
    if declared is not None and declared != extension_family:
        return None
    return extension_family


# ---------------------------------------------------------------------------
# Zip Bomb / Path Traversal / Active Content 방어 (DOCX/XLSX 공용).
# ---------------------------------------------------------------------------
class UnsafeContainerError(Exception):
    pass


class OutputLimitExceededError(Exception):
    """추출 도중 출력 길이 상한을 넘었다 - 전체를 끝까지 뽑은 뒤 검사하지
    않고 그 시점에서 즉시 중단한다."""


_ACTIVE_CONTENT_SUFFIXES = ("vbaproject.bin",)


def _contains_active_content(names) -> bool:
    return any(name.lower().endswith(suffix) for name in names for suffix in _ACTIVE_CONTENT_SUFFIXES)


def validate_zip_container(content: bytes, settings: Settings) -> None:
    """Central Directory의 선언된 크기'만' 믿지 않는다 - 선언된 크기를 위조한
    Archive를 잡기 위해 각 Entry를 실제로(다만 상한 안에서) 읽어 실제
    압축 해제 Byte 수를 선언된 값과 대조한다. 비율은 개별 Entry와 전체
    집계 양쪽 모두 검사한다(집계만 보면 "작은 악성 Entry 여러 개"를,
    개별만 보면 "큰 악성 Entry 하나가 다른 정상 Entry에 희석되는 경우"를
    놓칠 수 있다).

    비율 상한(기본 100:1)은 CLAUDE.md/M06 최초 설계가 제안했던 10:1이
    아니다 - 실제 python-docx가 만든 최소 DOCX 한 개조차 Word 자체의
    기본 `styles.xml`/`stylesWithEffects.xml` Boilerplate 때문에 약 24:1
    집계 비율을 갖는다(이 파일 작성 시점에 직접 측정: 문단 하나뿐인 문서가
    826,209 Byte 압축 해제 / 34,597 Byte 압축 ≈ 23.9:1). 10:1을 그대로
    적용하면 정상적인 실제 DOCX가 전부 거부된다 - 이는 "몰래 완화"가
    아니라 실측 근거로 재조정한 것이며, 실제 Zip Bomb(수백~수만 배)은
    여전히 100:1도 훨씬 넘는다(이 파일의 테스트 Fixture는 약 1000:1).
    """
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        infos = zf.infolist()
        if len(infos) > settings.max_container_entries:
            raise UnsafeContainerError("too many container entries")

        names = [info.filename for info in infos]
        if _contains_active_content(names):
            raise UnsafeContainerError("active content (macro) not permitted")

        declared_uncompressed_total = 0
        declared_compressed_total = 0
        for info in infos:
            name = info.filename
            if name.startswith("/") or ".." in name.replace("\\", "/").split("/"):
                raise UnsafeContainerError("path traversal entry")
            declared_uncompressed_total += info.file_size
            declared_compressed_total += info.compress_size
            if declared_uncompressed_total > settings.max_decompressed_bytes:
                raise UnsafeContainerError("decompressed size exceeds limit")
            compressed = max(info.compress_size, 1)
            if info.file_size / compressed > settings.max_compression_ratio:
                raise UnsafeContainerError("compression ratio exceeds limit")

        if declared_uncompressed_total > 0:
            aggregate_ratio = declared_uncompressed_total / max(declared_compressed_total, 1)
            if aggregate_ratio > settings.max_compression_ratio:
                raise UnsafeContainerError("compression ratio exceeds limit")

        # 실제 읽기 - 선언된 Central Directory 크기를 위조한 Archive를 잡는다.
        actual_total = 0
        for info in infos:
            with zf.open(info) as entry:
                read_for_entry = 0
                while True:
                    chunk = entry.read(65536)
                    if not chunk:
                        break
                    read_for_entry += len(chunk)
                    actual_total += len(chunk)
                    if actual_total > settings.max_decompressed_bytes:
                        raise UnsafeContainerError("decompressed size exceeds limit")
                if read_for_entry != info.file_size:
                    raise UnsafeContainerError("declared size does not match actual decompressed size")


# ---------------------------------------------------------------------------
# UTF-16 offset 변환 - Java String과 동일 단위로 위치를 표현한다(Non-BMP 문자는
# Java에서 Surrogate Pair 2개로 계산되므로 Python 코드포인트 오프셋과 다르다).
# ---------------------------------------------------------------------------
def _utf16_offset(text: str, codepoint_offset: int) -> int:
    return len(text[:codepoint_offset].encode("utf-16-le")) // 2


def _location(text: str, locator_type: str, locator_value: str, start_cp: int, end_cp: int) -> LocationDto:
    return LocationDto(
        locatorType=locator_type,
        locatorValue=locator_value,
        startOffset=_utf16_offset(text, start_cp),
        endOffset=_utf16_offset(text, end_cp),
    )


def _check_output_limit(current_length: int, settings: Settings) -> None:
    if current_length > settings.max_output_chars:
        raise OutputLimitExceededError("output length exceeds limit")


# ---------------------------------------------------------------------------
# 포맷별 Parser (Module-level 함수 - Multiprocessing Spawn이 Pickle할 수
# 있어야 한다). 각각 (parser_name, parser_version, full_text, locations)를
# 반환하거나 예외를 던진다. 출력 길이 상한은 전체를 다 모은 뒤가 아니라
# 누적하는 도중에 확인해 조기 중단한다.
# ---------------------------------------------------------------------------
def _parse_text(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    from charset_normalizer import from_bytes

    result = from_bytes(content).best()
    if result is None:
        raise ValueError("undecodable text content")
    text = str(result)
    _check_output_limit(len(text), settings)
    locations = [_location(text, "DOCUMENT", "1", 0, len(text))] if text else []
    return "charset-normalizer", "decode-only", text, locations


def _parse_pdf(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer

    full_text = ""
    locations: List[LocationDto] = []
    page_count = 0
    for page_layout in extract_pages(io.BytesIO(content)):
        page_count += 1
        if page_count > settings.max_pdf_pages:
            raise ValueError("page count exceeds limit")
        parts = [element.get_text() for element in page_layout if isinstance(element, LTTextContainer)]
        page_text = "".join(parts)
        start_cp = len(full_text)
        full_text += page_text
        _check_output_limit(len(full_text), settings)
        end_cp = len(full_text)
        if page_text:
            locations.append(_location(full_text, "PAGE", str(page_count), start_cp, end_cp))
    return "pdfminer.six", _module_version("pdfminer.six"), full_text, locations


def _parse_docx(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    """문단뿐 아니라 표(중첩 표 포함) 내용도 문서 순서대로 추출한다
    ({@code iter_inner_content}, python-docx>=1.1) - 표만 있는 문서가
    NO_TEXT로 잘못 분류되거나, 표 안의 사실이 조용히 누락되지 않는다."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(content))
    parts: List[str] = []
    length = 0

    def visit(container) -> None:
        nonlocal length
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph):
                if block.text:
                    parts.append(block.text)
                    length += len(block.text)
                    _check_output_limit(length, settings)
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        visit(cell)

    visit(document)
    full_text = "\n".join(parts)
    _check_output_limit(len(full_text), settings)
    locations = [_location(full_text, "DOCUMENT", "1", 0, len(full_text))] if full_text else []
    return "python-docx", _module_version("python-docx"), full_text, locations


def _parse_xlsx(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    """빈 첫 시트/선행-후행 공백에 관계없이 반환하는 `normalized_text`와
    `locations`가 항상 일치하도록, 빈 Segment를 join 대상에서 아예
    제외한다 - 나중에 `strip()`으로 잘라내지 않는다(그러면 이미 계산해 둔
    Offset이 어긋난다)."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        if len(sheet_names) > settings.max_xlsx_sheets:
            raise ValueError("sheet count exceeds limit")

        segments: List[Tuple[str, str]] = []
        running_length = 0
        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            rows_text: List[str] = []
            cell_count = 0
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    cell_count += 1
                    if cell_count > settings.max_xlsx_cells_per_sheet:
                        raise ValueError("cell count exceeds limit")
                    if value is not None:
                        rows_text.append(str(value))
            sheet_text = " ".join(rows_text)
            if sheet_text:
                running_length += len(sheet_text)
                _check_output_limit(running_length, settings)
                segments.append((sheet_name, sheet_text))

        full_text = "\n".join(text for _, text in segments)
        locations: List[LocationDto] = []
        cursor = 0
        for sheet_name, text in segments:
            start_cp = cursor
            end_cp = cursor + len(text)
            locations.append(_location(full_text, "SHEET_RANGE", sheet_name, start_cp, end_cp))
            cursor = end_cp + 1  # "\n" separator - harmless overshoot after the last segment.

        return "openpyxl", _module_version("openpyxl"), full_text, locations
    finally:
        workbook.close()


def _module_version(module_name: str) -> str:
    try:
        from importlib.metadata import version

        return version(module_name)
    except Exception:
        return "unknown"


def _sleep_forever_for_tests(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    """실제 실행 시간 Timeout+강제 종료(Kill)+자원 회수를 검증하기 위한
    테스트 전용 진입점이다. 어떤 실제 파일 형식도 이 Key로 분류되지 않으므로
    운영 경로로는 절대 도달할 수 없다."""
    import time as _time

    _time.sleep(10_000)
    return "test", "test", "unreachable", []


def _fail_immediately_for_tests(content: bytes, settings: Settings) -> Tuple[str, str, str, List[LocationDto]]:
    """자식 Process 안에서 일어나는 예외가 Parent에서 안전하게(원본 메시지
    노출 없이) FAILED로 번역되는지 결정적으로 검증하기 위한 테스트 전용
    진입점이다. 운영 경로로는 절대 도달할 수 없다."""
    raise RuntimeError("deliberate test failure - must never reach the HTTP response")


_PARSERS: dict = {
    _FAMILY_TEXT: _parse_text,
    _FAMILY_PDF: _parse_pdf,
    _FAMILY_DOCX: _parse_docx,
    _FAMILY_XLSX: _parse_xlsx,
    "_TEST_SLEEP_FOREVER": _sleep_forever_for_tests,
    "_TEST_FAIL_IMMEDIATELY": _fail_immediately_for_tests,
}


# ---------------------------------------------------------------------------
# 실제 Process 수준 Timeout 강제 + 결과 수신.
# ---------------------------------------------------------------------------
def _worker_entrypoint(content: bytes, file_name: str, declared_mime_type: Optional[str], settings: Settings,
        result_queue, family_override: Optional[str] = None) -> None:
    """분류/Container 검사/실제 Parsing을 전부 이 자식 Process 안에서
    수행한다 - Timeout과 동시성 상한이 신뢰할 수 없는 입력의 모든 처리
    단계에 적용되도록 한다(M06 후속 교정 - 이전에는 분류/검사가 이 경계
    밖 부모 Process에서 실행됐다). {@code family_override}는 테스트 전용
    통로다(`_TEST_SLEEP_FOREVER`/`_TEST_FAIL_IMMEDIATELY`) - 어떤 실제
    입력도 `classify_format`을 통해 이 값을 만들어낼 수 없으므로 운영
    경로에는 영향이 없다."""
    try:
        family = family_override if family_override is not None \
                else classify_format(content, file_name, declared_mime_type)
        if family is None:
            result_queue.put(("unsupported", "unsupported or mismatched format", None, None, None))
            return
        if family in (_FAMILY_DOCX, _FAMILY_XLSX):
            try:
                validate_zip_container(content, settings)
            except UnsafeContainerError as e:
                result_queue.put(("unsafe_container", str(e), None, None, None))
                return
            except zipfile.BadZipFile:
                result_queue.put(("unsafe_container", "malformed container", None, None, None))
                return
        parser_name, parser_version, text, locations = _PARSERS[family](content, settings)
        result_queue.put(("ok", parser_name, parser_version, text, [loc.model_dump() for loc in locations]))
    except OutputLimitExceededError:
        result_queue.put(("output_limit", "output length exceeds limit", None, None, None))
    except Exception as e:  # noqa: BLE001 - 자식 Process 경계를 넘겨야 하므로 넓게 잡는다.
        result_queue.put(("error", type(e).__name__, str(e), None, None))


def _terminate_and_reap(process: "mp.process.BaseProcess") -> None:
    """성공/예외/Timeout 세 경로 모두에서 호출된다({@code finally}) - Process
    Handle과 그 자원(메모리, OS Handle)을 항상 회수한다."""
    if process.is_alive():
        process.terminate()
        process.join(5)
        if process.is_alive():
            process.kill()
    process.join()


def _run_with_timeout(content: bytes, file_name: str, declared_mime_type: Optional[str],
        settings: Settings, family_override: Optional[str] = None) -> ParseResponseDto:
    """**중요(M06 후속 교정 - 실제 결함 수정):** 이전 구현은
    {@code process.join(timeout)}을 먼저 호출한 뒤에야 Queue를 읽었다.
    자식은 결과 전체를 Queue에 넣고, 그 Queue의 내부 Feeder Thread가
    Pipe에 다 쓸 때까지 기다린 뒤에야 종료한다 - 결과가 OS Pipe Buffer보다
    크면 자식은 "Parent가 읽어가기"를 기다리는데, Parent는 `join()`으로
    "자식이 끝나기"를 기다리므로 서로를 기다리며 교착한다(문서화된
    {@code multiprocessing.Queue}/{@code join} Deadlock 패턴이다 - 느린
    Parsing이 아니라 이 패턴 때문에 큰(정상) 결과가 거짓 Timeout으로
    보였다). 이제 Parent는 자식이 살아있는 동안 짧은 Timeout으로 Queue를
    반복 폴링해 결과가 도착하는 즉시 받는다 - {@code process.join(timeout)}을
    먼저 호출하지 않는다. 완료 여부 판단에 {@code queue.empty()}를 신뢰하지
    않는다(Race 가능) - 오직 `queue.get()`의 실제 반환/예외만 신뢰한다.
    """
    ctx = mp.get_context("spawn")
    result_queue = ctx.Queue()
    process = ctx.Process(target=_worker_entrypoint, args=(content, file_name, declared_mime_type, settings,
            result_queue, family_override), daemon=True)
    process.start()

    deadline = time.monotonic() + settings.parse_timeout_seconds
    message = None
    timed_out = False
    try:
        while message is None:
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                timed_out = True
                break
            try:
                message = result_queue.get(timeout=min(remaining, 0.25))
            except queue_module.Empty:
                if not process.is_alive():
                    # 자식이 결과를 넣기 전에 죽었을 수 있다 - 마지막으로 한 번
                    # Non-Blocking으로 더 확인한다(막 도착한 메시지와의 Race 대비).
                    try:
                        message = result_queue.get_nowait()
                    except queue_module.Empty:
                        pass
                    break
    finally:
        # 성공/예외/Timeout 모든 경로에서 Process/Queue 자원을 회수한다.
        _terminate_and_reap(process)
        result_queue.close()
        result_queue.join_thread()

    if message is None:
        if timed_out:
            raise TimeoutError("parser execution timed out")
        raise RuntimeError("parser process exited without a result")

    status, a, b, text, locations_raw = message
    if status == "ok":
        if not text or not text.strip():
            return ParseResponseDto(outcome=OUTCOME_NO_TEXT, reason="no extractable text")
        if len(text) > settings.max_output_chars:
            # 개별 Parser가 도중에 잡지 못했을 경우의 최종 방어선.
            return ParseResponseDto(outcome=OUTCOME_FAILED, reason="output length exceeds limit")
        locations = [LocationDto(**loc) for loc in locations_raw]
        return ParseResponseDto(outcome=OUTCOME_SUCCESS, parserName=a, parserVersion=b,
                normalizationVersion=NORMALIZATION_VERSION, normalizedText=text, locations=locations)
    if status == "unsupported":
        return ParseResponseDto(outcome=OUTCOME_UNSUPPORTED_FORMAT, reason=a)
    if status == "unsafe_container":
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason=f"unsafe container: {a}")
    if status == "output_limit":
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason="output length exceeds limit")
    # status == "error" - 자식의 원본 예외 메시지(b)는 절대 응답/로그에 노출하지 않는다.
    return ParseResponseDto(outcome=OUTCOME_FAILED, reason="parse error")


# ---------------------------------------------------------------------------
# Bounded Admission - 동시 활성 Parser 작업 수를 Non-Blocking으로 제한한다.
# 한도가 이미 찼으면 대기열에 쌓지 않고(Unbounded Waiting Queue 금지) 즉시
# 실패로 응답한다.
# ---------------------------------------------------------------------------
_concurrency_semaphore: Optional[threading.Semaphore] = None
_concurrency_lock = threading.Lock()


def _semaphore(settings: Settings) -> threading.Semaphore:
    global _concurrency_semaphore
    if _concurrency_semaphore is None:
        with _concurrency_lock:
            if _concurrency_semaphore is None:
                _concurrency_semaphore = threading.Semaphore(settings.max_concurrent_parses)
    return _concurrency_semaphore


def parse_document(content: bytes, file_name: str, declared_mime_type: Optional[str],
        settings: Settings) -> ParseResponseDto:
    if len(content) > settings.max_input_bytes:
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason="input size exceeds limit")

    semaphore = _semaphore(settings)
    if not semaphore.acquire(blocking=False):
        # 한도 초과 - 대기시키지 않고 즉시 실패로 응답한다(Bounded Admission).
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason="server at capacity")
    try:
        return _run_with_timeout(content, file_name, declared_mime_type, settings)
    except TimeoutError:
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason="parser execution timed out")
    except Exception:
        return ParseResponseDto(outcome=OUTCOME_FAILED, reason="parse error")
    finally:
        semaphore.release()
